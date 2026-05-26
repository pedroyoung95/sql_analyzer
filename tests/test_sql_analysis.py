import unittest
from pathlib import Path

from openpyxl import load_workbook

from sql_analyze_app.services.excel_export_service import exportResultsToExcel
from sql_analyze_app.services.sql_analysis_service import analyzeSourceTables, analyzeSql


class SqlAnalysisTest(unittest.TestCase):
    def testExampleSql(self):
        sqlText = """
        SELECT A.ORDER_ID,
               B.CUSTOMER_NAME
        FROM ORDERS A
        LEFT JOIN CUSTOMERS B
               ON A.CUSTOMER_ID = B.CUSTOMER_ID
        WHERE A.ORDER_DATE >= '2024-01-01'
          AND A.STATUS = 'COMPLETE';
        """

        results = analyzeSql(sqlText)

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0].tableName, "CUSTOMERS")
        self.assertEqual(results[0].conditionType, "JOIN")
        self.assertEqual(
            results[0].conditionText,
            "A.CUSTOMER_ID = B.CUSTOMER_ID",
        )
        self.assertEqual(results[1].tableName, "ORDERS")
        self.assertEqual(results[1].conditionType, "WHERE")
        self.assertIn("A.ORDER_DATE >= '2024-01-01'", results[1].conditionText)
        self.assertIn("A.STATUS = 'COMPLETE'", results[1].conditionText)

    def testCteAndSubquerySql(self):
        sqlText = """
        WITH RECENT_ORDERS AS (
            SELECT *
            FROM ORDERS O
            WHERE O.ORDER_DATE >= '2024-01-01'
        )
        SELECT R.ORDER_ID
        FROM RECENT_ORDERS R
        JOIN CUSTOMERS C ON R.CUSTOMER_ID = C.CUSTOMER_ID
        WHERE C.STATUS = 'Y'
        """

        results = analyzeSql(sqlText)
        whereRows = [result for result in results if result.conditionType == "WHERE"]
        tableNames = [result.tableName for result in results]

        self.assertIn("ORDERS", tableNames)
        self.assertIn("CUSTOMERS", tableNames)
        self.assertNotIn("RECENT_ORDERS", tableNames)
        self.assertTrue(any(row.tableName == "ORDERS" and "O.ORDER_DATE" in row.conditionText for row in whereRows))
        self.assertTrue(any(row.tableName == "CUSTOMERS" and "C.STATUS" in row.conditionText for row in whereRows))

    def testNestedWhereUsesCurrentQueryAlias(self):
        sqlText = """
        SELECT *
        FROM (
            SELECT O.ORDER_ID, O.CUSTOMER_ID
            FROM ORDERS O
            WHERE O.STATUS = 'COMPLETE'
        ) A
        JOIN CUSTOMERS C ON A.CUSTOMER_ID = C.CUSTOMER_ID
        WHERE C.STATUS = 'Y'
        """

        results = analyzeSql(sqlText)
        orderWhereRows = [
            result for result in results
            if result.tableName == "ORDERS" and result.conditionType == "WHERE"
        ]
        customerWhereRows = [
            result for result in results
            if result.tableName == "CUSTOMERS" and result.conditionType == "WHERE"
        ]

        self.assertTrue(any("O.STATUS = 'COMPLETE'" in row.conditionText for row in orderWhereRows))
        self.assertTrue(any("C.STATUS = 'Y'" in row.conditionText for row in customerWhereRows))
        self.assertFalse(any("C.STATUS = 'Y'" in row.conditionText for row in orderWhereRows))

    def testDerivedTableJoinConditionIsShownWithoutPhysicalTable(self):
        sqlText = """
        SELECT *
        FROM (
            SELECT A.EBELN, B.ACNG_DOC_NO, B.MTR
            FROM NTDWMART.DW_MM_BUY_DTL A
            INNER JOIN NTDWMART.DW_MM_BUY_HIST B
                ON A.EBELN = B.EBELN AND A.ITEM_NO = B.ITEM_NO
            WHERE B.HIST_KIND = 'E'
        ) A
        INNER JOIN (
            SELECT A.MTR PROD, B.ACNG_DOC_NO
            FROM NTDWMART.DW_MM_BUY_DTL A
            INNER JOIN NTDWMART.DW_MM_BUY_HIST B
                ON A.EBELN = B.EBELN AND A.ITEM_NO = B.ITEM_NO
            WHERE B.HIST_KIND = 'O'
        ) B
            ON A.ACNG_DOC_NO = B.ACNG_DOC_NO AND A.MTR = B.PROD
        """

        results = analyzeSql(sqlText)
        joinRows = [result for result in results if result.conditionType == "JOIN"]

        self.assertEqual(len(joinRows), 3)
        self.assertEqual(joinRows[0].tableName, "")
        self.assertEqual(joinRows[1].tableName, "NTDWMART.DW_MM_BUY_HIST")
        self.assertEqual(joinRows[2].tableName, "NTDWMART.DW_MM_BUY_HIST")
        self.assertTrue(any("A.ACNG_DOC_NO = B.ACNG_DOC_NO" in row.conditionText for row in joinRows))
        self.assertTrue(all(row.codeBlock.startswith("CODE BLOCK ") for row in joinRows))

    def testSourceTablesAreSeparatedFromConditions(self):
        sqlText = """
        SELECT *
        FROM NTDWMART.DW_MM_BUY_DTL A
        INNER JOIN NTDWMART.DW_MM_BUY_HIST B
            ON A.EBELN = B.EBELN
        WHERE B.HIST_KIND = 'E'
        """

        sourceTables = analyzeSourceTables(sqlText)
        results = analyzeSql(sqlText)

        self.assertEqual(
            [sourceTable.tableName for sourceTable in sourceTables],
            ["NTDWMART.DW_MM_BUY_DTL", "NTDWMART.DW_MM_BUY_HIST"],
        )
        self.assertEqual([result.conditionText for result in results], ["A.EBELN = B.EBELN", "B.HIST_KIND = 'E'"])

    def testExcelUsesSeparateTables(self):
        sqlText = """
        SELECT *
        FROM ORDERS A
        JOIN CUSTOMERS B ON A.CUSTOMER_ID = B.CUSTOMER_ID
        WHERE A.STATUS = 'Y'
        """

        results = analyzeSql(sqlText)
        sourceTables = analyzeSourceTables(sqlText)

        filePath = Path("excel_table_smoke.xlsx")
        exportResultsToExcel(results, str(filePath), sourceTables)
        workbook = load_workbook(filePath)
        sheet = workbook["SQL_ANALYSIS"]

        self.assertEqual(sheet["A1"].value, "SOURCE TABLES")
        self.assertEqual(sheet["D1"].value, "CONDITIONS")
        self.assertEqual([sheet["A2"].value, sheet["B2"].value], ["No", "Table Name"])
        self.assertEqual(
            [sheet["D2"].value, sheet["E2"].value, sheet["F2"].value, sheet["G2"].value],
            ["No", "Code Block", "Condition Type", "Condition Text"],
        )
        self.assertIn("SourceTableList", sheet.tables)
        self.assertIn("ConditionList", sheet.tables)
        workbook.close()

    def testExcelCanExportSourceTablesWithoutConditions(self):
        sqlText = "SELECT * FROM PRODUCTS"
        sourceTables = analyzeSourceTables(sqlText)

        filePath = Path("excel_source_only_smoke.xlsx")
        exportResultsToExcel([], str(filePath), sourceTables)
        workbook = load_workbook(filePath)
        sheet = workbook["SQL_ANALYSIS"]

        self.assertEqual(sheet["A1"].value, "SOURCE TABLES")
        self.assertEqual(sheet["B3"].value, "PRODUCTS")
        self.assertEqual(sheet["D1"].value, "CONDITIONS")
        workbook.close()

    def testIncludeTablesWithoutConditions(self):
        sqlText = "SELECT * FROM PRODUCTS"

        results = analyzeSql(sqlText, includeTablesWithoutConditions=True)

        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].tableName, "PRODUCTS")
        self.assertEqual(results[0].conditionType, "NONE")


if __name__ == "__main__":
    unittest.main()
