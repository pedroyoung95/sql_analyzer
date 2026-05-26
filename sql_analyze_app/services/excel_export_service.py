from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from sql_analyze_app.models.analysis_result import AnalysisResult
from sql_analyze_app.models.source_table_result import SourceTableResult


SOURCE_HEADERS = ["No", "Table Name"]
CONDITION_HEADERS = ["No", "Code Block", "Condition Type", "Condition Text"]


def exportResultsToExcel(
    results: list[AnalysisResult],
    filePath: str,
    sourceTables: list[SourceTableResult] | None = None,
) -> None:
    if not results and not sourceTables:
        raise ValueError("저장할 분석 결과가 없습니다.")

    path = Path(filePath)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SQL_ANALYSIS"

    sourceRows = sourceTables if sourceTables is not None else _sourceTablesFromResults(results)
    _writeSourceTable(sheet, sourceRows)
    _writeConditionTable(sheet, results)

    _styleSheet(sheet)
    workbook.save(path)


def _writeSourceTable(sheet, sourceTables: list[SourceTableResult]) -> None:
    sheet["A1"] = "SOURCE TABLES"
    for columnIndex, header in enumerate(SOURCE_HEADERS, start=1):
        sheet.cell(row=2, column=columnIndex, value=header)

    for rowIndex, sourceTable in enumerate(sourceTables, start=3):
        sheet.cell(row=rowIndex, column=1, value=sourceTable.no)
        sheet.cell(row=rowIndex, column=2, value=sourceTable.tableName)

    endRow = max(len(sourceTables) + 2, 3)
    if not sourceTables:
        sheet.cell(row=3, column=1, value="")
        sheet.cell(row=3, column=2, value="")
    _addExcelTable(sheet, "SourceTableList", f"A2:B{endRow}")


def _writeConditionTable(sheet, results: list[AnalysisResult]) -> None:
    startColumn = 4
    sheet.cell(row=1, column=startColumn, value="CONDITIONS")
    for offset, header in enumerate(CONDITION_HEADERS):
        sheet.cell(row=2, column=startColumn + offset, value=header)

    for rowIndex, result in enumerate(results, start=3):
        sheet.cell(row=rowIndex, column=startColumn, value=rowIndex - 2)
        sheet.cell(row=rowIndex, column=startColumn + 1, value=result.codeBlock)
        sheet.cell(row=rowIndex, column=startColumn + 2, value=result.conditionType)
        sheet.cell(row=rowIndex, column=startColumn + 3, value=result.conditionText)

    endRow = max(len(results) + 2, 3)
    _addExcelTable(sheet, "ConditionList", f"D2:G{endRow}")


def _addExcelTable(sheet, displayName: str, ref: str) -> None:
    table = Table(displayName=displayName, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    sheet.add_table(table)


def _sourceTablesFromResults(results: list[AnalysisResult]) -> list[SourceTableResult]:
    sourceTables = []
    seen = set()
    for result in results:
        key = result.tableName.upper()
        if key in seen:
            continue
        seen.add(key)
        sourceTables.append(SourceTableResult(len(sourceTables) + 1, result.tableName))
    return sourceTables


def _styleSheet(sheet) -> None:
    headerFill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    headerFont = Font(color="FFFFFF", bold=True)

    for cell in (sheet["A1"], sheet["D1"]):
        cell.font = Font(bold=True)

    for cell in list(sheet[2]):
        if cell.value:
            cell.fill = headerFill
            cell.font = headerFont
            cell.alignment = Alignment(horizontal="center")

    widths = [8, 34, 4, 8, 18, 18, 90]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=3):
        row[0].alignment = Alignment(horizontal="center")
        row[3].alignment = Alignment(horizontal="center")
        row[4].alignment = Alignment(horizontal="center")
        row[5].alignment = Alignment(horizontal="center")
        row[6].alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A3"
